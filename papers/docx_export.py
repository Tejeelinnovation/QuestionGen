"""
papers/docx_export.py

Professional Microsoft Word (.docx) Examination Paper and Answer Key Generator.
Supports:
- Institutional Exam Header (School name, Paper title, Version, Duration, Max marks).
- Instructions Banner and Candidate Identification Grid (Name, Roll No, Date, Invigilator).
- Standard Pedagogical Section Grouping (Section A: MCQs, Section B: Short, Section C: Long).
- Configurable Blank Writing Lines:
  - Universal baseline (e.g. N blank lines per mark).
  - Per-question individual overrides (teacher can give 10 lines to a 3-mark derivation and 4 to a 3-mark definition).
- Synchronized Teacher Answer Key with correct choices, model solutions, and topic tags.
"""

from __future__ import annotations

import io
import logging
from typing import Any, Dict, List, Optional, Union

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from papers.models import PaperVersion
from papers.publishing import ExamSection, FormattedExamQuestion, PaperPublishingService

logger = logging.getLogger(__name__)


class ExamPaperDocxService:
    """
    Generates publication-ready Word (.docx) exam papers and answer keys.
    """

    COLOR_PRIMARY = RGBColor(30, 41, 59)      # Dark Navy #1E293B
    COLOR_INDIGO = RGBColor(67, 56, 202)      # Indigo #4338CA
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Grey #64748B
    COLOR_LINE = RGBColor(160, 174, 192)      # Ruled line color #A0AEC0
    COLOR_CORRECT = RGBColor(21, 128, 61)     # Green #15803D

    def __init__(self):
        self.publishing_service = PaperPublishingService()

    def generate_student_paper(
        self,
        version: PaperVersion,
        lines_per_mark: int = 2,
        question_line_overrides: Optional[Dict[Union[str, int], int]] = None,
        line_style: str = "dotted",
    ) -> io.BytesIO:
        """
        Generate a student question paper Word (.docx) document with blank writing lines.
        """
        overrides = {str(k): int(v) for k, v in (question_line_overrides or {}).items()}
        sections = self.publishing_service.compile_sections(version)

        doc = docx.Document()
        self._configure_page_layout(doc)

        # 1. Institutional Header
        self._add_institutional_header(doc, version)

        # 2. Candidate Identification Box
        self._add_candidate_identification_table(doc)

        # 3. Candidate Instructions (if provided)
        self._add_instructions_box(doc, version)

        # 4. Sections & Questions Body
        for sec in sections:
            if not sec.questions:
                continue

            self._add_section_header(doc, sec)

            for q in sec.questions:
                # Calculate required blank lines for this question
                q_id_str = str(q.question_id)
                if q_id_str in overrides:
                    blank_lines = max(0, overrides[q_id_str])
                else:
                    if q.question_type == "MCQ":
                        blank_lines = 0
                    else:
                        blank_lines = max(0, int(round(q.marks * lines_per_mark)))

                self._add_student_question(doc, q, blank_lines=blank_lines, line_style=line_style)

        # 5. End of Paper Marker
        self._add_end_marker(doc, is_answer_key=False)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    def generate_answer_key(self, version: PaperVersion) -> io.BytesIO:
        """
        Generate a synchronized Teacher Answer Key and Marking Scheme Word (.docx) document.
        """
        sections = self.publishing_service.compile_sections(version)

        doc = docx.Document()
        self._configure_page_layout(doc)

        # 1. Answer Key Title Header
        self._add_answer_key_header(doc, version)

        # 2. Sections & Solutions Body
        for sec in sections:
            if not sec.questions:
                continue

            self._add_section_header(doc, sec, is_answer_key=True)

            for q in sec.questions:
                self._add_teacher_solution(doc, q)

        # 3. End Marker
        self._add_end_marker(doc, is_answer_key=True)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    # ---------------------------------------------------------------------------
    # Internal Document Styling & Builder Helpers
    # ---------------------------------------------------------------------------

    def _configure_page_layout(self, doc: docx.Document) -> None:
        """Set standard page margins and default styles."""
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Normal text style
        normal_style = doc.styles["Normal"]
        normal_style.font.name = "Calibri"
        normal_style.font.size = Pt(11)
        normal_style.font.color.rgb = self.COLOR_PRIMARY

    def _add_institutional_header(self, doc: docx.Document, version: PaperVersion) -> None:
        """Add school name, examination title, metadata row, and subjects."""
        paper = version.paper
        school_name = getattr(paper.school, "name", "ACADEMIC INSTITUTION").upper() if paper.school else "ACADEMIC INSTITUTION"

        # School Name
        p_school = doc.add_paragraph()
        p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_school.paragraph_format.space_before = Pt(0)
        p_school.paragraph_format.space_after = Pt(2)
        r_school = p_school.add_run(school_name)
        r_school.bold = True
        r_school.font.size = Pt(13)
        r_school.font.color.rgb = self.COLOR_PRIMARY

        # Paper Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(4)
        r_title = p_title.add_run(paper.title.upper())
        r_title.bold = True
        r_title.font.size = Pt(16)
        r_title.font.color.rgb = self.COLOR_INDIGO

        # Metadata Row
        p_meta = doc.add_paragraph()
        p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_meta.paragraph_format.space_before = Pt(0)
        p_meta.paragraph_format.space_after = Pt(4)

        items = [
            f"VERSION: {version.version_label}",
            f"TOTAL MARKS: {float(version.total_marks or 0.0):g}",
        ]
        if paper.duration_minutes:
            items.append(f"TIME ALLOWED: {paper.duration_minutes} MIN")
        if version.question_snapshot:
            items.append(f"QUESTIONS: {len(version.question_snapshot)}")

        r_meta = p_meta.add_run("   |   ".join(items))
        r_meta.bold = True
        r_meta.font.size = Pt(9.5)
        r_meta.font.color.rgb = self.COLOR_MUTED

        # Subjects (if any)
        if paper.subjects:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(8)
            r_sub = p_sub.add_run(f"SUBJECTS: {' • '.join(paper.subjects)}")
            r_sub.font.size = Pt(9)
            r_sub.font.color.rgb = self.COLOR_MUTED

    def _add_answer_key_header(self, doc: docx.Document, version: PaperVersion) -> None:
        """Header for Teacher Answer Key and Solution Scheme."""
        paper = version.paper
        school_name = getattr(paper.school, "name", "ACADEMIC INSTITUTION").upper() if paper.school else "ACADEMIC INSTITUTION"

        p_school = doc.add_paragraph()
        p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_school.paragraph_format.space_after = Pt(2)
        r_school = p_school.add_run(school_name)
        r_school.bold = True
        r_school.font.size = Pt(12)
        r_school.font.color.rgb = self.COLOR_MUTED

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(3)
        r_title = p_title.add_run(f"{paper.title.upper()} — ANSWER KEY")
        r_title.bold = True
        r_title.font.size = Pt(16)
        r_title.font.color.rgb = self.COLOR_CORRECT

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(10)
        r_sub = p_sub.add_run(f"OFFICIAL MARKING SCHEME & MODEL SOLUTIONS  |  VERSION {version.version_label}")
        r_sub.bold = True
        r_sub.font.size = Pt(9.5)
        r_sub.font.color.rgb = self.COLOR_PRIMARY

    def _add_candidate_identification_table(self, doc: docx.Document) -> None:
        """2x2 bordered grid for student identification details."""
        table = doc.add_table(rows=2, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        col_widths = [Inches(3.5), Inches(3.5)]
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = col_widths[i]

        cells = table.rows[0].cells
        self._format_meta_cell(cells[0], "STUDENT NAME:", "_____________________________________")
        self._format_meta_cell(cells[1], "EXAMINATION DATE:", "_________________________")

        cells_2 = table.rows[1].cells
        self._format_meta_cell(cells_2[0], "ROLL / ADMISSION ID:", "_____________________________________")
        self._format_meta_cell(cells_2[1], "INVIGILATOR SIGNATURE:", "_________________________")

        self._set_table_borders(table)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def _format_meta_cell(self, cell, label: str, fill_line: str) -> None:
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r_label = p.add_run(f"{label} ")
        r_label.bold = True
        r_label.font.size = Pt(8.5)
        r_label.font.color.rgb = self.COLOR_PRIMARY
        r_line = p.add_run(fill_line)
        r_line.font.size = Pt(8.5)
        r_line.font.color.rgb = self.COLOR_MUTED

    def _set_table_borders(self, table) -> None:
        """Apply subtle borders around a docx table."""
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
            f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>\n'
            f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def _add_instructions_box(self, doc: docx.Document, version: PaperVersion) -> None:
        instructions = (version.paper.instructions or "").strip()
        if not instructions:
            return

        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        cell.width = Inches(7.0)

        # Background color #F8FAFC
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
        tcPr.append(shd)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r_prefix = p.add_run("GENERAL INSTRUCTIONS: ")
        r_prefix.bold = True
        r_prefix.font.size = Pt(8.5)
        r_prefix.font.color.rgb = self.COLOR_PRIMARY

        r_text = p.add_run(instructions)
        r_text.italic = True
        r_text.font.size = Pt(8.5)
        r_text.font.color.rgb = self.COLOR_PRIMARY

        self._set_table_borders(table)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def _add_section_header(self, doc: docx.Document, sec: ExamSection, is_answer_key: bool = False) -> None:
        """Section banner with marks breakdown."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True

        r_title = p.add_run(sec.title.upper())
        r_title.bold = True
        r_title.font.size = Pt(11)
        r_title.font.color.rgb = self.COLOR_INDIGO if not is_answer_key else self.COLOR_CORRECT

        r_meta = p.add_run(f"  [{len(sec.questions)} Questions • Total {sec.total_marks:g} Marks]")
        r_meta.bold = True
        r_meta.font.size = Pt(9)
        r_meta.font.color.rgb = self.COLOR_MUTED

        if sec.instructions and not is_answer_key:
            p_inst = doc.add_paragraph()
            p_inst.paragraph_format.space_before = Pt(0)
            p_inst.paragraph_format.space_after = Pt(4)
            p_inst.paragraph_format.keep_with_next = True
            r_inst = p_inst.add_run(sec.instructions)
            r_inst.italic = True
            r_inst.font.size = Pt(8.5)
            r_inst.font.color.rgb = self.COLOR_MUTED

    def _add_student_question(
        self,
        doc: docx.Document,
        q: FormattedExamQuestion,
        blank_lines: int = 0,
        line_style: str = "dotted",
    ) -> None:
        """Render a single question and its options / blank answer lines."""
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(3)
        p_q.paragraph_format.keep_with_next = True

        # Question Number
        r_num = p_q.add_run(f"Q{q.global_question_number}.  ")
        r_num.bold = True
        r_num.font.size = Pt(10)
        r_num.font.color.rgb = self.COLOR_PRIMARY

        # Question Text
        r_text = p_q.add_run(q.question_text)
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = self.COLOR_PRIMARY

        # Marks tag
        marks_str = f"  [{q.marks:g} {'Mark' if q.marks == 1 else 'Marks'}]"
        r_marks = p_q.add_run(marks_str)
        r_marks.bold = True
        r_marks.font.size = Pt(9)
        r_marks.font.color.rgb = self.COLOR_INDIGO

        # Render Question Diagram if present
        if q.diagram:
            fig_id = q.diagram.get("figure_id")
            if fig_id:
                try:
                    from papers.services.pdf_harvester import PdfInspectionService
                    img_bytes = PdfInspectionService.get_figure_bytes(fig_id)

                    # Tier 2: Check persistent local media storage
                    if not img_bytes:
                        try:
                            from django.conf import settings
                            local_path = os.path.join(settings.MEDIA_ROOT, "diagrams", f"{fig_id}.png")
                            if os.path.exists(local_path):
                                with open(local_path, "rb") as lf:
                                    img_bytes = lf.read()
                        except Exception:
                            pass

                    # Tier 3: Check base64 data URI in image_url
                    if not img_bytes and str(q.diagram.get("image_url", "")).startswith("data:image/"):
                        try:
                            b64_part = q.diagram["image_url"].split(",", 1)[-1]
                            img_bytes = base64.b64decode(b64_part)
                        except Exception:
                            pass

                    # Tier 4: Check thumbnail base64
                    if not img_bytes and q.diagram.get("thumbnail_b64"):
                        try:
                            img_bytes = base64.b64decode(q.diagram["thumbnail_b64"])
                        except Exception:
                            pass

                    # Tier 5: Google Drive download
                    if not img_bytes:
                        try:
                            from content_ingestion.storage.drive_client import GoogleDriveClient
                            drive_client = GoogleDriveClient()
                            drive_file_id = q.diagram.get("drive_file_id")
                            if drive_file_id and drive_client.is_configured:
                                img_bytes = drive_client.download_bytes(drive_file_id)
                        except Exception:
                            pass

                    if img_bytes:
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(4)
                        p_img.paragraph_format.space_after = Pt(4)
                        p_img.paragraph_format.keep_with_next = True
                        r_img = p_img.add_run()
                        r_img.add_picture(io.BytesIO(img_bytes), width=Inches(3.8))

                        # Caption
                        cap = q.diagram.get("title")
                        if cap:
                            p_cap = doc.add_paragraph()
                            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p_cap.paragraph_format.space_before = Pt(0)
                            p_cap.paragraph_format.space_after = Pt(4)
                            p_cap.paragraph_format.keep_with_next = True
                            r_cap = p_cap.add_run(cap)
                            r_cap.italic = True
                            r_cap.font.size = Pt(8.5)
                            r_cap.font.color.rgb = self.COLOR_MUTED
                    else:
                        # Graceful callout if image bytes could not be located
                        cap = q.diagram.get("title") or "Refer to diagram attached with question"
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r_cap = p_cap.add_run(f"[Diagram: {cap}]")
                        r_cap.italic = True
                        r_cap.font.size = Pt(9)
                        r_cap.font.color.rgb = self.COLOR_INDIGO
                except Exception as img_err:
                    logger.debug(f"Could not embed diagram image into DOCX: {img_err}")

        # MCQ Options
        if q.options and isinstance(q.options, dict):
            opts_table = doc.add_table(rows=0, cols=2)
            opts_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            opts_table.autofit = False

            col_widths = [Inches(3.3), Inches(3.3)]
            keys = sorted(q.options.keys())

            # Pair up keys (A, B), (C, D)
            for i in range(0, len(keys), 2):
                row = opts_table.add_row()
                for j, cell in enumerate(row.cells):
                    cell.width = col_widths[j]
                    idx = i + j
                    if idx < len(keys):
                        k = keys[idx]
                        v = q.options[k]
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.space_before = Pt(1)
                        cp.paragraph_format.space_after = Pt(2)
                        r_k = cp.add_run(f"({k}) ")
                        r_k.bold = True
                        r_k.font.size = Pt(9.5)
                        r_v = cp.add_run(str(v))
                        r_v.font.size = Pt(9.5)

        # Ruled Blank Answer Lines
        if blank_lines > 0:
            for _ in range(blank_lines):
                p_line = doc.add_paragraph()
                p_line.paragraph_format.space_before = Pt(1)
                p_line.paragraph_format.space_after = Pt(2)
                p_line.paragraph_format.left_indent = Inches(0.2)

                if line_style == "solid":
                    line_chars = "_________________________________________________________________________________"
                else:  # dotted
                    line_chars = ".  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  .  ."

                r_line = p_line.add_run(line_chars)
                r_line.font.size = Pt(8)
                r_line.font.color.rgb = self.COLOR_LINE

    def _add_teacher_solution(self, doc: docx.Document, q: FormattedExamQuestion) -> None:
        """Render question with correct answer and model solutions for teacher."""
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(8)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.keep_with_next = True

        r_num = p_q.add_run(f"Q{q.global_question_number}.  ")
        r_num.bold = True
        r_num.font.size = Pt(10)

        r_text = p_q.add_run(q.question_text)
        r_text.font.size = Pt(10)

        r_marks = p_q.add_run(f"  [{q.marks:g} Marks]")
        r_marks.bold = True
        r_marks.font.size = Pt(9)
        r_marks.font.color.rgb = self.COLOR_MUTED

        # Correct Answer block
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Inches(0.25)
        p_ans.paragraph_format.space_before = Pt(2)
        p_ans.paragraph_format.space_after = Pt(2)

        r_ans_label = p_ans.add_run("CORRECT ANSWER: ")
        r_ans_label.bold = True
        r_ans_label.font.size = Pt(9.5)
        r_ans_label.font.color.rgb = self.COLOR_CORRECT

        if q.question_type == "MCQ" and q.options:
            correct_key = q.correct_answer.strip().upper()
            opt_text = q.options.get(correct_key, "")
            r_ans_val = p_ans.add_run(f"Option ({correct_key}) — {opt_text}" if opt_text else f"Option ({correct_key})")
        else:
            r_ans_val = p_ans.add_run(q.correct_answer or "Descriptive solution evaluation per rubric.")

        r_ans_val.bold = True
        r_ans_val.font.size = Pt(9.5)
        r_ans_val.font.color.rgb = self.COLOR_CORRECT

        # Source / Topic info
        if q.topic_name or q.source_reference:
            p_ref = doc.add_paragraph()
            p_ref.paragraph_format.left_indent = Inches(0.25)
            p_ref.paragraph_format.space_before = Pt(0)
            p_ref.paragraph_format.space_after = Pt(4)
            ref_items = []
            if q.topic_name:
                ref_items.append(f"Topic: {q.topic_name}")
            if q.source_reference:
                ref_items.append(f"Reference: {q.source_reference}")
            r_ref = p_ref.add_run("   |   ".join(ref_items))
            r_ref.font.size = Pt(8.5)
            r_ref.font.color.rgb = self.COLOR_MUTED

    def _add_end_marker(self, doc: docx.Document, is_answer_key: bool = False) -> None:
        """Document terminator footer."""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(10)

        marker_text = (
            "*** END OF EXAMINATION PAPER • ALL THE BEST ***"
            if not is_answer_key
            else "*** END OF TEACHER ANSWER KEY & MARKING SCHEME ***"
        )
        r = p.add_run(marker_text)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = self.COLOR_MUTED
