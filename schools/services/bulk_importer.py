"""
Bulk Import Service for Students and Teachers.
Implements header validation, format sanitization, deterministic normalization,
duplicate checking, transactional/atomic capacity enforcement, and comprehensive audit reporting.
"""

from __future__ import annotations

import io
import re
from typing import Any
import openpyxl
from django.contrib.auth import get_user_model
from django.db import transaction

from core.audit import log_action
from schools.models import ClassSection, ClassSubjectTeacher, School
from users.models import Capability, UserCapability
from .normalization import AcademicClassNormalizer

User = get_user_model()

INDIAN_PHONE_REGEX = re.compile(r"^\+91[6-9]\d{9}$")


def clean_indian_phone(raw_phone: Any) -> str | None:
    """
    Cleans and validates an Indian mobile number into canonical '+91XXXXXXXXXX' format.
    Accepts 10 digits, digits with +91, leading 0, spaces, hyphens, or floats from Excel.
    """
    if raw_phone is None:
        return None

    # Handle numeric float/int from openpyxl
    if isinstance(raw_phone, (int, float)):
        raw_phone = str(int(raw_phone))

    cleaned = re.sub(r"[\s\-\(\)\.]", "", str(raw_phone).strip())
    if not cleaned:
        return None

    # Strip leading +
    if cleaned.startswith("+"):
        cleaned = cleaned[1:]

    # If starts with 0 and length is 11, strip 0
    if cleaned.startswith("0") and len(cleaned) == 11:
        cleaned = cleaned[1:]

    # If 10 digits, prepend 91
    if len(cleaned) == 10 and cleaned[0] in "6789":
        cleaned = "91" + cleaned

    # Prepend +
    formatted = "+" + cleaned
    if INDIAN_PHONE_REGEX.match(formatted):
        return formatted

    return None


def sanitize_username_string(name: str) -> str:
    """Sanitizes text into safe alphanumeric characters for usernames."""
    cleaned = re.sub(r"[^a-zA-Z0-9]", "", name.lower())
    return cleaned or "user"


class BulkImporter:
    """
    Bulk Importer for processing Excel workbooks for student and teacher cohorts.
    """

    @classmethod
    def import_students_from_excel(
        cls, file_content: bytes, school_id: int, creator_user: User
    ) -> dict[str, Any]:
        """
        Processes a Student Excel workbook with atomic capacity enforcement.
        """
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            ws = wb.active
        except Exception as e:
            return {
                "error": f"Invalid Excel workbook format: {str(e)}",
                "summary": {"total_rows": 0, "created_count": 0, "over_limit_count": 0, "duplicate_count": 0, "invalid_count": 0},
                "created": [],
                "over_limit": [],
                "duplicates": [],
                "invalid": [{"row_number": 0, "error": f"Failed to parse Excel file: {str(e)}"}],
            }

        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            return {
                "error": "The uploaded workbook does not contain any data rows.",
                "summary": {"total_rows": 0, "created_count": 0, "over_limit_count": 0, "duplicate_count": 0, "invalid_count": 0},
                "created": [],
                "over_limit": [],
                "duplicates": [],
                "invalid": [{"row_number": 1, "error": "No data rows found below header."}],
            }

        # Header matching
        raw_headers = [str(cell or "").strip() for cell in rows[0]]
        header_map: dict[str, int] = {}
        for idx, h in enumerate(raw_headers):
            h_lower = h.lower()
            if "student" in h_lower and "name" in h_lower:
                header_map["name"] = idx
            elif "name" in h_lower and "name" not in header_map:
                header_map["name"] = idx
            elif "gr" in h_lower:
                header_map["gr_number"] = idx
            elif "roll" in h_lower:
                header_map["roll_number"] = idx
            elif "standard" in h_lower or "class" in h_lower:
                header_map["standard"] = idx
            elif "div" in h_lower or "sec" in h_lower:
                header_map["division"] = idx
            elif "phone" in h_lower or "mobile" in h_lower:
                header_map["phone"] = idx

        required_keys = ["name", "gr_number", "standard", "division", "phone"]
        missing_keys = [k for k in required_keys if k not in header_map]
        if missing_keys:
            return {
                "error": f"Missing required columns in workbook: {', '.join(missing_keys)}",
                "summary": {"total_rows": len(rows) - 1, "created_count": 0, "over_limit_count": 0, "duplicate_count": 0, "invalid_count": len(rows) - 1},
                "created": [],
                "over_limit": [],
                "duplicates": [],
                "invalid": [{"row_number": 1, "error": f"Missing required columns: {', '.join(missing_keys)}. Expected headers: Student Name, GR Number, Roll Number, Standard, Division, Phone Number."}],
            }

        # Atomic execution with database row lock on School
        with transaction.atomic():
            school = School.objects.select_for_update().get(id=school_id)
            current_student_count = school.users.filter(role="Student").count()
            configured_limit = school.max_students
            remaining_capacity = max(0, configured_limit - current_student_count)

            report: dict[str, Any] = {
                "summary": {
                    "total_rows": len(rows) - 1,
                    "created_count": 0,
                    "over_limit_count": 0,
                    "duplicate_count": 0,
                    "invalid_count": 0,
                    "configured_limit": configured_limit,
                    "current_usage_before": current_student_count,
                    "current_usage_after": current_student_count,
                    "remaining_capacity_before": remaining_capacity,
                    "remaining_capacity_after": remaining_capacity,
                },
                "created": [],
                "over_limit": [],
                "duplicates": [],
                "invalid": [],
            }

            # In-file tracking
            seen_gr_numbers: set[str] = set()
            seen_phones: set[str] = set()

            # Existing records in DB for this school / system
            existing_grs = set(
                school.users.filter(role="Student")
                .exclude(gr_number="")
                .values_list("gr_number", flat=True)
            )
            existing_phones = set(User.objects.values_list("mobile_number", flat=True))

            # Fetch capability models
            student_caps = Capability.objects.filter(name__in=["ATTEMPT_TEST", "VIEW_OWN_RESULT"])

            for row_idx, row_values in enumerate(rows[1:], start=2):
                # Check if entire row is empty
                if not any(row_values):
                    continue

                raw_name = str(row_values[header_map["name"]] or "").strip() if "name" in header_map else ""
                raw_gr = str(row_values[header_map["gr_number"]] or "").strip() if "gr_number" in header_map else ""
                raw_roll = (
                    str(row_values[header_map["roll_number"]] or "").strip()
                    if "roll_number" in header_map and header_map["roll_number"] < len(row_values)
                    else ""
                )
                raw_standard = (
                    row_values[header_map["standard"]] if "standard" in header_map else None
                )
                raw_division = (
                    str(row_values[header_map["division"]] or "").strip()
                    if "division" in header_map
                    else None
                )
                raw_phone = (
                    row_values[header_map["phone"]] if "phone" in header_map else None
                )

                # 1. Validation Checks
                if not raw_name:
                    report["invalid"].append({
                        "row_number": row_idx,
                        "name": raw_name or "Unknown",
                        "gr_number": raw_gr,
                        "error": "Student Name is required.",
                    })
                    continue

                if not raw_gr:
                    report["invalid"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "gr_number": "",
                        "error": "GR Number is required as an organization identifier.",
                    })
                    continue

                norm_standard = AcademicClassNormalizer.normalize_standard(raw_standard)
                if not norm_standard:
                    report["invalid"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "gr_number": raw_gr,
                        "error": f"Invalid Standard/Class '{raw_standard}'. Expected standard 1-12 (e.g. Standard 10, 10th, X).",
                    })
                    continue

                norm_division = AcademicClassNormalizer.normalize_division(raw_division)
                if not norm_division:
                    report["invalid"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "gr_number": raw_gr,
                        "error": f"Invalid Division/Section '{raw_division}'. Expected single letter division (e.g. A-J).",
                    })
                    continue

                norm_phone = clean_indian_phone(raw_phone)
                if not norm_phone:
                    report["invalid"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "gr_number": raw_gr,
                        "error": f"Invalid Phone Number '{raw_phone}'. Must be a valid 10-digit Indian mobile number (+91).",
                    })
                    continue

                # 2. Duplicate Checks
                if raw_gr in seen_gr_numbers or raw_gr in existing_grs:
                    report["duplicates"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "gr_number": raw_gr,
                        "duplicate_field": "GR Number",
                        "error": f"Student with GR Number '{raw_gr}' already exists in this school.",
                    })
                    continue

                if norm_phone in seen_phones or norm_phone in existing_phones:
                    report["duplicates"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "gr_number": raw_gr,
                        "duplicate_field": "Phone Number",
                        "error": f"Mobile number '{norm_phone}' is already registered in the system.",
                    })
                    continue

                # 3. Capacity Enforcement
                if remaining_capacity <= 0:
                    report["over_limit"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "gr_number": raw_gr,
                        "standard": norm_standard,
                        "division": norm_division,
                        "error": f"School student capacity limit of {configured_limit} has been reached. Account was not created.",
                    })
                    continue

                # 4. Create Account
                seen_gr_numbers.add(raw_gr)
                seen_phones.add(norm_phone)

                # Parse first / last name
                name_parts = raw_name.split(maxsplit=1)
                first_name = name_parts[0]
                last_name = name_parts[1] if len(name_parts) > 1 else ""

                # Deterministic safe username
                clean_gr = sanitize_username_string(raw_gr)
                base_username = f"st_{clean_gr}"
                username = base_username
                counter = 1
                while User.objects.filter(username=username).exists():
                    username = f"{base_username}_{counter}"
                    counter += 1

                email = f"{clean_gr}@student.school{school.id}.edu"
                if User.objects.filter(email=email).exists():
                    email = f"{clean_gr}_{counter}@student.school{school.id}.edu"

                # Resolve or create ClassSection
                class_sec, _ = ClassSection.objects.get_or_create(
                    school=school,
                    standard=norm_standard,
                    section=norm_division,
                    defaults={"max_students": 40},
                )

                user = User(
                    username=username,
                    email=email,
                    first_name=first_name,
                    last_name=last_name,
                    mobile_number=norm_phone,
                    role="Student",
                    school=school,
                    created_by=creator_user,
                    class_section=class_sec,
                    gr_number=raw_gr,
                    roll_number=raw_roll,
                )
                user.set_password("Welcome@123")
                user.save()

                # Grant capabilities
                for cap in student_caps:
                    UserCapability.objects.create(user=user, capability=cap, granted_by=creator_user)

                remaining_capacity -= 1
                report["created"].append({
                    "row_number": row_idx,
                    "id": user.id,
                    "name": raw_name,
                    "username": username,
                    "gr_number": raw_gr,
                    "roll_number": raw_roll,
                    "standard": norm_standard,
                    "division": norm_division,
                    "class_name": f"Class {norm_standard}-{norm_division}",
                    "mobile_number": norm_phone,
                })

            # Update final summaries
            report["summary"]["created_count"] = len(report["created"])
            report["summary"]["over_limit_count"] = len(report["over_limit"])
            report["summary"]["duplicate_count"] = len(report["duplicates"])
            report["summary"]["invalid_count"] = len(report["invalid"])
            report["summary"]["current_usage_after"] = current_student_count + len(report["created"])
            report["summary"]["remaining_capacity_after"] = remaining_capacity

            log_action(
                creator_user,
                "bulk_import.students",
                school,
                metadata={
                    "created": report["summary"]["created_count"],
                    "over_limit": report["summary"]["over_limit_count"],
                    "duplicates": report["summary"]["duplicate_count"],
                    "invalid": report["summary"]["invalid_count"],
                },
            )

            return report

    @classmethod
    def import_teachers_from_excel(
        cls, file_content: bytes, school_id: int, creator_user: User
    ) -> dict[str, Any]:
        """
        Processes a Teacher Excel workbook with atomic capacity enforcement.
        """
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            ws = wb.active
        except Exception as e:
            return {
                "error": f"Invalid Excel workbook format: {str(e)}",
                "summary": {"total_rows": 0, "created_count": 0, "over_limit_count": 0, "duplicate_count": 0, "invalid_count": 0},
                "created": [],
                "over_limit": [],
                "duplicates": [],
                "invalid": [{"row_number": 0, "error": f"Failed to parse Excel file: {str(e)}"}],
            }

        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            return {
                "error": "The uploaded workbook does not contain any data rows.",
                "summary": {"total_rows": 0, "created_count": 0, "over_limit_count": 0, "duplicate_count": 0, "invalid_count": 0},
                "created": [],
                "over_limit": [],
                "duplicates": [],
                "invalid": [{"row_number": 1, "error": "No data rows found below header."}],
            }

        # Header matching
        raw_headers = [str(cell or "").strip() for cell in rows[0]]
        header_map: dict[str, int] = {}
        for idx, h in enumerate(raw_headers):
            h_lower = h.lower()
            if "teacher" in h_lower and "name" in h_lower:
                header_map["name"] = idx
            elif "name" in h_lower and "name" not in header_map:
                header_map["name"] = idx
            elif "phone" in h_lower or "mobile" in h_lower:
                header_map["phone"] = idx
            elif "subject" in h_lower:
                header_map["subject"] = idx
            elif "class teacher" in h_lower or "in-charge" in h_lower:
                header_map["class_teacher"] = idx

        required_keys = ["name", "phone", "subject"]
        missing_keys = [k for k in required_keys if k not in header_map]
        if missing_keys:
            return {
                "error": f"Missing required columns in workbook: {', '.join(missing_keys)}",
                "summary": {"total_rows": len(rows) - 1, "created_count": 0, "over_limit_count": 0, "duplicate_count": 0, "invalid_count": len(rows) - 1},
                "created": [],
                "over_limit": [],
                "duplicates": [],
                "invalid": [{"row_number": 1, "error": f"Missing required columns: {', '.join(missing_keys)}. Expected headers: Teacher Name, Teacher Mobile Number, Subject, Class Teacher."}],
            }

        with transaction.atomic():
            school = School.objects.select_for_update().get(id=school_id)
            current_teacher_count = school.users.filter(role="Teacher").count()
            configured_limit = school.max_teachers
            remaining_capacity = max(0, configured_limit - current_teacher_count)

            report: dict[str, Any] = {
                "summary": {
                    "total_rows": len(rows) - 1,
                    "created_count": 0,
                    "over_limit_count": 0,
                    "duplicate_count": 0,
                    "invalid_count": 0,
                    "configured_limit": configured_limit,
                    "current_usage_before": current_teacher_count,
                    "current_usage_after": current_teacher_count,
                    "remaining_capacity_before": remaining_capacity,
                    "remaining_capacity_after": remaining_capacity,
                },
                "created": [],
                "over_limit": [],
                "duplicates": [],
                "invalid": [],
            }

            seen_phones: set[str] = set()
            existing_phones = set(User.objects.values_list("mobile_number", flat=True))

            teacher_caps = Capability.objects.filter(
                name__in=["CREATE_PAPER", "ASSIGN_TEST", "GENERATE_SELECT_QUESTIONS"]
            )

            for row_idx, row_values in enumerate(rows[1:], start=2):
                if not any(row_values):
                    continue

                raw_name = str(row_values[header_map["name"]] or "").strip() if "name" in header_map else ""
                raw_phone = row_values[header_map["phone"]] if "phone" in header_map else None
                raw_subject = str(row_values[header_map["subject"]] or "").strip() if "subject" in header_map else ""
                raw_ct = (
                    str(row_values[header_map["class_teacher"]] or "").strip()
                    if "class_teacher" in header_map and header_map["class_teacher"] < len(row_values)
                    else ""
                )

                # 1. Validation
                if not raw_name:
                    report["invalid"].append({
                        "row_number": row_idx,
                        "name": raw_name or "Unknown",
                        "error": "Teacher Name is required.",
                    })
                    continue

                norm_phone = clean_indian_phone(raw_phone)
                if not norm_phone:
                    report["invalid"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "error": f"Invalid Mobile Number '{raw_phone}'. Must be a valid 10-digit Indian mobile number (+91).",
                    })
                    continue

                if not raw_subject:
                    report["invalid"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "error": "Subject is required.",
                    })
                    continue

                # 2. Duplicate Check
                if norm_phone in seen_phones or norm_phone in existing_phones:
                    report["duplicates"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "duplicate_field": "Mobile Number",
                        "error": f"Teacher with mobile number '{norm_phone}' is already registered in the system.",
                    })
                    continue

                # 3. Capacity Enforcement
                if remaining_capacity <= 0:
                    report["over_limit"].append({
                        "row_number": row_idx,
                        "name": raw_name,
                        "subject": raw_subject,
                        "error": f"School teacher capacity limit of {configured_limit} has been reached. Account was not created.",
                    })
                    continue

                # 4. Parse Class Teacher assignment
                parsed_ct = AcademicClassNormalizer.parse_class_teacher_assignment(raw_ct)

                # 5. Create Teacher Account
                seen_phones.add(norm_phone)

                name_parts = raw_name.split(maxsplit=1)
                first_name = name_parts[0]
                last_name = name_parts[1] if len(name_parts) > 1 else ""

                clean_name = sanitize_username_string(first_name)
                phone_suffix = norm_phone[-4:]
                base_username = f"tch_{clean_name}_{phone_suffix}"
                username = base_username
                counter = 1
                while User.objects.filter(username=username).exists():
                    username = f"{base_username}_{counter}"
                    counter += 1

                email = f"{username}@school{school.id}.edu"
                if User.objects.filter(email=email).exists():
                    email = f"{username}_{counter}@school{school.id}.edu"

                user = User(
                    username=username,
                    email=email,
                    first_name=first_name,
                    last_name=last_name,
                    mobile_number=norm_phone,
                    role="Teacher",
                    school=school,
                    created_by=creator_user,
                    primary_subject=raw_subject,
                )
                user.set_password("Welcome@123")
                user.save()

                for cap in teacher_caps:
                    UserCapability.objects.create(user=user, capability=cap, granted_by=creator_user)

                # Assign Class Teacher and subject if parsed
                assigned_ct_info = None
                if parsed_ct.standard and parsed_ct.section:
                    sec, _ = ClassSection.objects.get_or_create(
                        school=school,
                        standard=parsed_ct.standard,
                        section=parsed_ct.section,
                        defaults={"max_students": 40},
                    )
                    sec.class_teacher = user
                    sec.class_teacher_subject = raw_subject
                    sec.save()

                    # Also map subject
                    ClassSubjectTeacher.objects.get_or_create(
                        class_section=sec,
                        subject=raw_subject,
                        defaults={"teacher": user},
                    )
                    assigned_ct_info = f"Class {parsed_ct.standard}-{parsed_ct.section} ({raw_subject})"

                remaining_capacity -= 1
                report["created"].append({
                    "row_number": row_idx,
                    "id": user.id,
                    "name": raw_name,
                    "username": username,
                    "mobile_number": norm_phone,
                    "subject": raw_subject,
                    "class_teacher": assigned_ct_info or "None",
                })

            report["summary"]["created_count"] = len(report["created"])
            report["summary"]["over_limit_count"] = len(report["over_limit"])
            report["summary"]["duplicate_count"] = len(report["duplicates"])
            report["summary"]["invalid_count"] = len(report["invalid"])
            report["summary"]["current_usage_after"] = current_teacher_count + len(report["created"])
            report["summary"]["remaining_capacity_after"] = remaining_capacity

            log_action(
                creator_user,
                "bulk_import.teachers",
                school,
                metadata={
                    "created": report["summary"]["created_count"],
                    "over_limit": report["summary"]["over_limit_count"],
                    "duplicates": report["summary"]["duplicate_count"],
                    "invalid": report["summary"]["invalid_count"],
                },
            )

            return report
